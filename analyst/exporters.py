"""Экспорт артефактов задачи: Markdown, PDF, DOCX, OpenAPI YAML, JSON, Mermaid, PlantUML, Jira."""

from __future__ import annotations

import io
import json
import os
import re
import time
from pathlib import Path
from typing import Any

import yaml

from .models import ApiSpec, RequirementsDoc
from .renderers import artifact_markdown, diagram_sources, md_risks, openapi_yaml

ORDER = [
    ("discovery", "Контекст и интервью"),
    ("requirements", "Требования"),
    ("review", "Ревью требований"),
    ("architecture", "Архитектура"),
    ("architecture_review", "Review архитектурного решения"),
    ("api", "API"),
    ("sequence", "Sequence Diagram"),
    ("activity", "Activity / BPMN"),
    ("erd", "ERD"),
    ("test_cases", "Test Cases"),
    ("final_review", "Финальная проверка"),
]

FORMATS = {
    "md": ("text/markdown; charset=utf-8", "md"),
    "pdf": ("application/pdf", "pdf"),
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
    "json": ("application/json", "json"),
    "openapi": ("application/yaml", "yaml"),
    "mermaid": ("text/plain; charset=utf-8", "mmd"),
    "plantuml": ("text/plain; charset=utf-8", "puml"),
    "jira": ("text/markdown; charset=utf-8", "jira.md"),
}


class ExportError(ValueError):
    pass


def _content(artifacts: dict[str, dict[str, Any]], kind: str) -> dict[str, Any] | None:
    art = artifacts.get(kind)
    return art["content"] if art else None


def risks_markdown(artifacts: dict[str, dict[str, Any]]) -> str:
    from .models import Risk

    risks = []
    review = _content(artifacts, "review")
    if review:
        risks += review["review"].get("risks", [])
    arch = _content(artifacts, "architecture")
    if arch:
        risks += arch.get("risks", [])
    return md_risks([Risk.model_validate(r) for r in risks]) if risks else "Риски не выявлены."


def markdown_report(project: dict[str, Any], task: dict[str, Any], artifacts: dict[str, dict[str, Any]]) -> str:
    parts = [f"# {task['title']}", f"**Проект:** {project['name']}  ",
             f"**Дата:** {time.strftime('%d.%m.%Y %H:%M')}  ",
             f"**Статус задачи:** {task['stage']}", "## Постановка задачи", task["problem"]]
    for kind, title in ORDER:
        art = artifacts.get(kind)
        if not art:
            continue
        status = " (опубликовано)" if art.get("status") == "published" else ""
        parts.append(f"## {title} — v{art['version']}{status}")
        parts.append(artifact_markdown(kind, art["content"]))
    if "review" in artifacts or "architecture" in artifacts:
        parts += ["## Риски", risks_markdown(artifacts)]
    parts.append("---\n_Сгенерировано System Analyst Agent. Артефакты требуют проверки аналитиком._")
    return "\n\n".join(parts)


def diagrams(artifacts: dict[str, dict[str, Any]], fmt: str) -> str:
    blocks = []
    for kind in ("sequence", "activity", "erd", "architecture"):
        content = _content(artifacts, kind)
        if content:
            src = diagram_sources(kind, content)[fmt]
            title = content.get("title") or "Component diagram"
            comment = "%%" if fmt == "mermaid" else "'"
            blocks.append(f"{comment} {kind}: {title}\n{src}")
    if not blocks:
        raise ExportError("В задаче ещё нет диаграмм")
    return "\n\n".join(blocks)


def jira_issue(task: dict[str, Any], artifacts: dict[str, dict[str, Any]]) -> str:
    """Черновик Jira Issue. Создание задачи в Jira — этап 2 и только после подтверждения."""
    req = _content(artifacts, "requirements")
    if not req:
        raise ExportError("Сначала сформируйте требования")
    doc = RequirementsDoc.model_validate(req)
    api = _content(artifacts, "api")
    arch = _content(artifacts, "architecture")
    parts = [f"**Summary:** {doc.title}", "h2. Description", doc.summary,
             "h2. Business Requirements", "\n".join(f"* {r.id}: {r.text}" for r in doc.business_requirements),
             "h2. Functional Requirements", "\n".join(f"* {r.id}: {r.text}" for r in doc.functional_requirements)]
    if api:
        spec = ApiSpec.model_validate(api["spec"])
        parts += ["h2. API Contract", "\n".join(f"* {e.method} {e.path} — {e.summary}" for e in spec.endpoints),
                  "OpenAPI: см. вложение openapi.yaml"]
    parts += ["h2. Acceptance Criteria",
              "\n".join(f"* {a.id}: Given {a.given} When {a.when} Then {a.then}" for a in doc.acceptance_criteria)]
    deps = []
    if arch:
        deps = [c["name"] for c in arch["components"] if not c["is_new"]] + arch.get("external_systems", [])
    parts += ["h2. Dependencies", "\n".join(f"* {d}" for d in deps) or "—",
              "h2. Technical Notes",
              "\n".join(f"* {n.id} ({n.category}): {n.text} — {n.metric}" for n in doc.non_functional_requirements)
              + ("\n" + "\n".join(f"* Open: {q}" for q in doc.open_issues) if doc.open_issues else "")]
    return "\n\n".join(parts)


# --- PDF ------------------------------------------------------------------------------------

FONT_CANDIDATES = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",
    "/Library/Fonts/Arial Unicode.ttf",
    "C:/Windows/Fonts/arial.ttf",
]


def _font(bold: bool = False) -> str | None:
    env = os.environ.get("ANALYST_PDF_FONT_BOLD" if bold else "ANALYST_PDF_FONT")
    candidates = ([env] if env else []) + [
        c.replace("DejaVuSans.ttf", "DejaVuSans-Bold.ttf") if bold else c for c in FONT_CANDIDATES]
    return next((c for c in candidates if c and Path(c).exists()), None)


PDF_REPLACEMENTS = {"⛔": "[!]", "✓": "+"}


def _plain(text: str) -> str:
    for char, repl in PDF_REPLACEMENTS.items():
        text = text.replace(char, repl)
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"(?<!\w)_(.+?)_(?!\w)", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return re.sub(r"</?(details|summary|small|br/?)>", "", text)


def markdown_to_pdf(markdown: str) -> bytes:
    from fpdf import FPDF

    regular = _font()
    if not regular:
        raise ExportError("Для PDF нужен TTF-шрифт с кириллицей: задайте ANALYST_PDF_FONT (например, DejaVuSans.ttf)")
    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_font("Main", "", regular)
    pdf.add_font("Main", "B", _font(bold=True) or regular)
    pdf.add_page()
    width = pdf.w - pdf.l_margin - pdf.r_margin

    def write(text: str, size: float = 10, style: str = "", h: float = 5) -> None:
        pdf.set_font("Main", style, size)
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(width, h, text or " ")

    in_code = False
    for line in markdown.splitlines():
        if line.startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            write(line, 8, h=4)
            continue
        heading = re.match(r"^(#{1,4})\s+(.*)", line)
        if heading:
            pdf.ln(2)
            write(_plain(heading.group(2)), {1: 16, 2: 14, 3: 12, 4: 11}[len(heading.group(1))], "B", 7)
        elif line.startswith("|---") or line.startswith("|---|"):
            continue
        elif line.startswith("|"):
            cells = [c.strip() for c in line.strip("|").split(" | ")]
            write(_plain("  ·  ".join(cells).replace("<br>", "; ")), 8, h=4)
        elif line.strip():
            write(_plain(line))
        else:
            pdf.ln(2)
    return bytes(pdf.output())


# --- DOCX ------------------------------------------------------------------------------------

def markdown_to_docx(markdown: str) -> bytes:
    import docx
    from docx.shared import Pt

    document = docx.Document()
    table = None
    in_code = False
    for line in markdown.splitlines():
        if line.startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            p = document.add_paragraph()
            run = p.add_run(line)
            run.font.name = "Courier New"
            run.font.size = Pt(8)
            continue
        if line.startswith("|"):
            if line.startswith("|---"):
                continue
            cells = [_plain(c.strip()).replace("<br>", "\n") for c in line.strip().strip("|").split(" | ")]
            if table is None:
                table = document.add_table(rows=0, cols=len(cells))
                table.style = "Table Grid"
            row = table.add_row().cells
            for i, value in enumerate(cells[: len(row)]):
                row[i].text = value
            continue
        table = None
        heading = re.match(r"^(#{1,4})\s+(.*)", line)
        if heading:
            document.add_heading(_plain(heading.group(2)), level=min(len(heading.group(1)), 4))
        elif re.match(r"^\s*[-*] ", line):
            document.add_paragraph(_plain(line.strip()[2:]), style="List Bullet")
        elif re.match(r"^\s*\d+\. ", line):
            document.add_paragraph(_plain(re.sub(r"^\s*\d+\.\s", "", line)), style="List Number")
        elif line.strip():
            document.add_paragraph(_plain(line))
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def export(fmt: str, project: dict[str, Any], task: dict[str, Any],
           artifacts: dict[str, dict[str, Any]]) -> tuple[bytes, str, str]:
    """Возвращает (данные, media type, имя файла)."""
    if fmt not in FORMATS:
        raise ExportError(f"Неизвестный формат {fmt}: {', '.join(FORMATS)}")
    if not artifacts:
        raise ExportError("В задаче ещё нет артефактов")
    media_type, ext = FORMATS[fmt]
    slug = re.sub(r"[^\w-]+", "-", task["title"].lower(), flags=re.UNICODE).strip("-")[:60] or "task"
    if fmt == "md":
        data = markdown_report(project, task, artifacts).encode()
    elif fmt == "pdf":
        data = markdown_to_pdf(markdown_report(project, task, artifacts))
    elif fmt == "docx":
        data = markdown_to_docx(markdown_report(project, task, artifacts))
    elif fmt == "json":
        payload = {"project": project["name"], "task": {k: task[k] for k in ("id", "title", "problem", "stage")},
                   "artifacts": {k: {"version": a["version"], "status": a.get("status"), "content": a["content"]}
                                 for k, a in artifacts.items()}}
        data = json.dumps(payload, ensure_ascii=False, indent=2).encode()
    elif fmt == "openapi":
        api = _content(artifacts, "api")
        if not api:
            raise ExportError("В задаче ещё нет API-контракта")
        data = openapi_yaml(ApiSpec.model_validate(api["spec"])).encode()
        slug = "openapi"
    elif fmt in ("mermaid", "plantuml"):
        data = diagrams(artifacts, fmt).encode()
    else:
        data = jira_issue(task, artifacts).encode()
    return data, media_type, f"{slug}.{ext}"


def openapi_is_valid(text: str) -> bool:
    doc = yaml.safe_load(text)
    return isinstance(doc, dict) and str(doc.get("openapi", "")).startswith("3.") and "paths" in doc
