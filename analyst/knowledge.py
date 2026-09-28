"""База знаний проекта: разбор документов, категории, чанки и поиск.

Поиск — BM25 по чанкам одного проекта. Интерфейс `Retriever` позволяет
заменить его на embeddings + векторную БД без изменений в агентах.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import re
import threading
import xml.etree.ElementTree as ET
from collections import Counter
from dataclasses import dataclass
from pathlib import PurePath
from typing import Any, Protocol

import yaml

from .storage import Store

MAX_UPLOAD_BYTES = 20 * 1024 * 1024
CHUNK_CHARS = 1200
CHUNK_OVERLAP = 150

CATEGORIES = ("Requirements", "API", "Architecture", "Database", "Business Rules", "Glossary", "Documentation")

SUPPORTED = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".markdown": "text/markdown",
    ".json": "application/json",
    ".yaml": "application/yaml",
    ".yml": "application/yaml",
    ".xml": "application/xml",
    ".csv": "text/csv",
}

HTTP_METHODS = ("get", "post", "put", "patch", "delete", "head", "options")


class DocumentError(ValueError):
    pass


@dataclass
class ParsedDocument:
    text: str
    category: str
    meta: dict[str, Any]
    media_type: str
    sha256: str


# --- разбор -------------------------------------------------------------------------

def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join(f"[стр. {i}]\n{page.extract_text() or ''}" for i, page in enumerate(reader.pages, 1))


def _docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = []
    for p in document.paragraphs:
        if not p.text.strip():
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        level = re.search(r"heading (\d)", style)
        parts.append(("#" * int(level.group(1)) + " " if level else "") + p.text)
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _xlsx(data: bytes) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"## Лист: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(cells):
                parts.append(" | ".join(cells))
    wb.close()
    return "\n".join(parts)


def _xml(data: bytes) -> str:
    if b"<!DOCTYPE" in data[:4096].upper() or b"<!ENTITY" in data.upper():
        raise DocumentError("XML с DOCTYPE/ENTITY не поддерживается из соображений безопасности")
    root = ET.fromstring(data)
    lines = []

    def walk(node: ET.Element, path: str) -> None:
        tag = node.tag.split("}")[-1]
        here = f"{path}/{tag}"
        attrs = " ".join(f'{k.split("}")[-1]}="{v}"' for k, v in node.attrib.items())
        text = (node.text or "").strip()
        if attrs or text:
            lines.append(f"{here} {attrs} {text}".strip())
        for child in node:
            walk(child, here)

    walk(root, "")
    return "\n".join(lines)


def extract_openapi_endpoints(spec: Any) -> list[str]:
    """«POST /api/v1/refunds — Create refund» для каждой операции OpenAPI/Swagger."""
    if not isinstance(spec, dict) or not ("openapi" in spec or "swagger" in spec):
        return []
    endpoints = []
    for path, item in (spec.get("paths") or {}).items():
        if not isinstance(item, dict):
            continue
        for method, op in item.items():
            if method.lower() in HTTP_METHODS:
                summary = (op.get("summary") or op.get("operationId") or "") if isinstance(op, dict) else ""
                endpoints.append(f"{method.upper()} {path}" + (f" — {summary}" if summary else ""))
    return endpoints


def _structured(text: str, ext: str) -> tuple[str, dict[str, Any]]:
    try:
        data = json.loads(text) if ext == ".json" else yaml.safe_load(text)
    except (json.JSONDecodeError, yaml.YAMLError):
        return text, {}
    endpoints = extract_openapi_endpoints(data)
    meta: dict[str, Any] = {}
    if endpoints:
        meta["openapi"] = True
        meta["endpoints"] = endpoints
        info = data.get("info") or {}
        meta["api_title"] = str(info.get("title", ""))
        text = "Операции API:\n" + "\n".join(f"- {e}" for e in endpoints) + "\n\n" + text
    return text, meta


def categorize(filename: str, text: str, meta: dict[str, Any]) -> str:
    if meta.get("openapi"):
        return "API"
    name = filename.lower()
    sample = text[:5000].lower()
    rules = [
        ("Glossary", ("glossary", "глоссар", "термин")),
        ("Business Rules", ("business-rule", "business_rule", "rules", "правил")),
        ("Database", ("database", "schema", "erd", "ddl", "create table", "таблиц", " бд", "db_", "db-")),
        ("Architecture", ("architecture", "архитектур", "c4", "component", "компонент", "deployment")),
        ("API", ("api", "endpoint", "swagger", "openapi", "эндпоинт")),
        ("Requirements", ("requirement", "требован", "тз", "srs", "brd", "user stor", "acceptance")),
    ]
    for category, keys in rules:
        if any(k in name for k in keys):
            return category
    for category, keys in rules:
        if sum(sample.count(k) for k in keys) >= 3:
            return category
    return "Documentation"


def parse_document(filename: str, data: bytes, category: str | None = None) -> ParsedDocument:
    if len(data) > MAX_UPLOAD_BYTES:
        raise DocumentError(f"Файл больше {MAX_UPLOAD_BYTES // (1024 * 1024)} МБ")
    ext = PurePath(filename).suffix.lower()
    if ext not in SUPPORTED:
        raise DocumentError(f"Формат {ext or '(без расширения)'} не поддерживается: {', '.join(SUPPORTED)}")
    meta: dict[str, Any] = {}
    try:
        if ext == ".pdf":
            text = _pdf(data)
        elif ext == ".docx":
            text = _docx(data)
        elif ext == ".xlsx":
            text = _xlsx(data)
        elif ext == ".xml":
            text = _xml(data)
        elif ext == ".csv":
            text = "\n".join(" | ".join(row) for row in csv.reader(io.StringIO(_decode(data))))
        else:
            text = _decode(data)
            if ext in (".json", ".yaml", ".yml"):
                text, meta = _structured(text, ext)
    except DocumentError:
        raise
    except Exception as e:  # повреждённый файл — понятная ошибка вместо 500
        raise DocumentError(f"Не удалось прочитать {filename}: {e}") from e
    text = text.strip()
    if not text:
        raise DocumentError(f"В {filename} не найден текст (скан без OCR?)")
    if category and category not in CATEGORIES:
        raise DocumentError(f"Неизвестная категория {category}: {', '.join(CATEGORIES)}")
    return ParsedDocument(
        text=text,
        category=category or categorize(filename, text, meta),
        meta=meta,
        media_type=SUPPORTED[ext],
        sha256=hashlib.sha256(data).hexdigest(),
    )


def chunk_text(text: str, size: int = CHUNK_CHARS, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """Режет по абзацам, длинные абзацы — по размеру, с перекрытием."""
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    chunks: list[str] = []
    current = ""
    for p in paragraphs:
        while len(p) > size:
            if current:
                chunks.append(current)
                current = ""
            chunks.append(p[:size])
            p = p[size - overlap:]
        if len(current) + len(p) + 2 > size and current:
            chunks.append(current)
            current = current[-overlap:] + "\n\n" + p if overlap else p
        else:
            current = f"{current}\n\n{p}" if current else p
    if current:
        chunks.append(current)
    return chunks


# --- поиск ------------------------------------------------------------------------------

_WORD = re.compile(r"[\w/{}.-]+", re.UNICODE)


def tokenize(text: str) -> list[str]:
    tokens = []
    for raw in _WORD.findall(text.lower()):
        raw = raw.strip(".-")
        parts = [raw] + [p for p in re.split(r"[/{}.-]+", raw) if p != raw]
        for word in parts:
            if len(word) < 2:
                continue
            # Грубый стемминг для русских и английских словоформ: «возврата» → «возвра».
            tokens.append(word[:6] if word.isalpha() and len(word) > 6 else word)
    return tokens


@dataclass
class Hit:
    document_id: str
    filename: str
    category: str
    text: str
    score: float


class Retriever(Protocol):
    def search(self, project_id: str, query: str, k: int = 6) -> list[Hit]: ...


class BM25Retriever:
    """BM25 по чанкам проекта. Индекс строится лениво и сбрасывается при
    изменении набора документов; индексы разных проектов не пересекаются."""

    def __init__(self, store: Store, k1: float = 1.5, b: float = 0.75):
        self.store = store
        self.k1, self.b = k1, b
        self._cache: dict[str, tuple[str, dict[str, Any]]] = {}
        self._lock = threading.Lock()

    def _index(self, project_id: str) -> dict[str, Any]:
        version = self.store.knowledge_version(project_id)
        with self._lock:
            cached = self._cache.get(project_id)
            if cached and cached[0] == version:
                return cached[1]
        chunks = self.store.chunks(project_id)
        docs = [Counter(tokenize(c["text"] + " " + c["filename"])) for c in chunks]
        df: Counter[str] = Counter()
        for tf in docs:
            df.update(tf.keys())
        lengths = [sum(tf.values()) for tf in docs]
        index = {"chunks": chunks, "tf": docs, "df": df, "len": lengths,
                 "avg": (sum(lengths) / len(lengths)) if lengths else 0.0}
        with self._lock:
            self._cache[project_id] = (version, index)
        return index

    def search(self, project_id: str, query: str, k: int = 6) -> list[Hit]:
        index = self._index(project_id)
        n = len(index["chunks"])
        if not n:
            return []
        terms = set(tokenize(query))
        scored = []
        for i, tf in enumerate(index["tf"]):
            score = 0.0
            for term in terms:
                if term not in tf:
                    continue
                idf = math.log(1 + (n - index["df"][term] + 0.5) / (index["df"][term] + 0.5))
                freq = tf[term]
                norm = self.k1 * (1 - self.b + self.b * index["len"][i] / (index["avg"] or 1))
                score += idf * freq * (self.k1 + 1) / (freq + norm)
            if score > 0:
                scored.append((score, i))
        scored.sort(reverse=True)
        hits = []
        for score, i in scored[:k]:
            c = index["chunks"][i]
            hits.append(Hit(c["document_id"], c["filename"], c["category"], c["text"], round(score, 3)))
        return hits


def existing_endpoints(store: Store, project_id: str) -> list[str]:
    """Все операции из загруженных OpenAPI-спецификаций и контекста проекта."""
    found: list[str] = []
    for doc in store.documents(project_id):
        found.extend(doc["meta"].get("endpoints", []))
    manual = store.get_project(project_id)["context"].get("existing_apis", "")
    found.extend(line.strip("-• ").strip() for line in str(manual).splitlines() if line.strip())
    return list(dict.fromkeys(found))


def knowledge_tree(store: Store, project_id: str) -> dict[str, list[dict[str, Any]]]:
    tree: dict[str, list[dict[str, Any]]] = {c: [] for c in CATEGORIES}
    for doc in store.documents(project_id):
        tree.setdefault(doc["category"], []).append(
            {"id": doc["id"], "filename": doc["filename"], "size": doc["size"]}
        )
    return tree
