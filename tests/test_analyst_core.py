from __future__ import annotations

import io
import json
from types import SimpleNamespace

import pytest
import yaml

from analyst.agents.orchestrator import endpoint_conflicts
from analyst.checklist import quick_check, score, traceability
from analyst.exporters import export, openapi_is_valid
from analyst.knowledge import (
    BM25Retriever, DocumentError, chunk_text, existing_endpoints, parse_document, tokenize,
)
from analyst.llm import ClaudeLLM, LLMRefusal, strict_schema
from analyst.models import ApiSpec, DiscoveryResult, ReviewResult, SequenceDiagram
from analyst.renderers import (
    activity_mermaid, activity_plantuml, artifact_markdown, component_mermaid, erd_mermaid, erd_plantuml,
    fields_schema, openapi_document, sequence_mermaid, sequence_plantuml,
)
from analyst.security import Cipher
from analyst.storage import Store

from analyst_fakes import ACTIVITY, API, ARCHITECTURE, ERD, REQUIREMENTS, REVIEW, SEQUENCE, TESTS

OPENAPI = """
openapi: 3.0.0
info: {title: Payments, version: '1'}
paths:
  /api/v1/payments/{id}:
    get: {summary: Get payment}
  /api/v1/refunds:
    post: {operationId: createRefund}
"""


# --- база знаний ------------------------------------------------------------------------

def test_openapi_upload_extracts_endpoints_and_category():
    doc = parse_document("existing-api.yaml", OPENAPI.encode())
    assert doc.category == "API"
    assert doc.meta["endpoints"] == ["GET /api/v1/payments/{id} — Get payment", "POST /api/v1/refunds — createRefund"]
    assert "Операции API" in doc.text


def test_categorize_by_name_and_content():
    assert parse_document("business-rules.md", "Правило 1".encode()).category == "Business Rules"
    assert parse_document("glossary.txt", b"term").category == "Glossary"
    assert parse_document("notes.txt", "Архитектура: компонент A, компонент B, архитектура".encode()).category \
        == "Architecture"
    assert parse_document("x.txt", b"hello").category == "Documentation"
    assert parse_document("x.txt", b"hello", category="Database").category == "Database"


def test_docx_xlsx_xml_and_cp1251():
    import docx
    import openpyxl

    d = docx.Document()
    d.add_heading("Бизнес-правила", 1)
    d.add_paragraph("Возврат не больше суммы платежа")
    buf = io.BytesIO()
    d.save(buf)
    parsed = parse_document("rules.docx", buf.getvalue())
    assert "# Бизнес-правила" in parsed.text and "Возврат" in parsed.text

    wb = openpyxl.Workbook()
    wb.active.append(["table", "column", "type"])
    wb.active.append(["refund", "id", "uuid"])
    buf = io.BytesIO()
    wb.save(buf)
    assert "refund | id | uuid" in parse_document("database.xlsx", buf.getvalue()).text

    assert "/root/item" in parse_document("a.xml", b"<root><item id='1'>x</item></root>").text
    with pytest.raises(DocumentError):
        parse_document("a.xml", b"<!DOCTYPE x [<!ENTITY a 'b'>]><root>&a;</root>")
    assert "Возврат" in parse_document("a.txt", "Возврат".encode("cp1251")).text


def test_rejects_unknown_format_and_empty():
    with pytest.raises(DocumentError):
        parse_document("a.exe", b"x")
    with pytest.raises(DocumentError):
        parse_document("a.txt", b"   ")


def test_chunking_and_tokenize():
    text = "\n\n".join(f"Абзац {i} " + "слово " * 50 for i in range(20))
    chunks = chunk_text(text, size=600, overlap=50)
    assert len(chunks) > 5 and all(len(c) <= 700 for c in chunks)
    assert tokenize("Возврата возврат") == ["возвра", "возвра"]


def test_retrieval_is_isolated_by_project():
    store = Store()
    user = store.create_user("u", "h")
    p1 = store.create_project("A", "", {}, user["id"])
    p2 = store.create_project("B", "", {}, user["id"])
    for pid, text in ((p1["id"], "Возврат платежа на карту клиента"), (p2["id"], "Секретный возврат проекта B")):
        store.add_document(pid, filename="d.txt", media_type="text/plain", category="Documentation", size=1,
                           sha256=pid, content=text, meta={}, user_id=user["id"], chunks=[text])
    retriever = BM25Retriever(store)
    hits = retriever.search(p1["id"], "возврат")
    assert len(hits) == 1 and "Секретный" not in hits[0].text
    # Индекс обновляется после загрузки нового документа.
    store.add_document(p1["id"], filename="e.txt", media_type="text/plain", category="Documentation", size=1,
                       sha256="x", content="возврат частичный", meta={}, user_id=user["id"],
                       chunks=["возврат частичный"])
    assert len(retriever.search(p1["id"], "возврат")) == 2


def test_encryption_at_rest(tmp_path):
    key = Cipher.generate_key()
    store = Store(tmp_path / "db.sqlite", Cipher(key))
    user = store.create_user("u", "h")
    p = store.create_project("A", "", {}, user["id"])
    store.add_document(p["id"], filename="d.txt", media_type="text/plain", category="Documentation", size=1,
                       sha256="s", content="секрет", meta={}, user_id=user["id"], chunks=["секрет"])
    raw = (tmp_path / "db.sqlite").read_bytes()
    assert "секрет".encode() not in raw
    assert store.chunks(p["id"])[0]["text"] == "секрет"


def test_existing_endpoints_from_docs_and_context():
    store = Store()
    user = store.create_user("u", "h")
    p = store.create_project("A", "", {"existing_apis": "- GET /api/v1/customers"}, user["id"])
    parsed = parse_document("api.yaml", OPENAPI.encode())
    store.add_document(p["id"], filename="api.yaml", media_type=parsed.media_type, category=parsed.category,
                       size=1, sha256="s", content=parsed.text, meta=parsed.meta, user_id=user["id"], chunks=["x"])
    assert existing_endpoints(store, p["id"])[-1] == "GET /api/v1/customers"
    assert len(existing_endpoints(store, p["id"])) == 3


# --- проверка требований ----------------------------------------------------------------------

def test_quick_check_flags_underspecified_requirement():
    result = quick_check("Сервис должен отправлять запрос на возврат платежа.")
    assert not result["detailed_enough"]
    for aspect in ("Timeout / retry", "Idempotency", "Logging", "Monitoring"):
        assert aspect in result["missing_aspects"]
    assert "Что происходит при повторном запросе?" in result["questions"]


def test_quick_check_finds_vague_wording():
    result = quick_check("Система должна работать быстро и удобно.")
    assert result["vague"][0]["term"] == "быстро"


def test_score_penalizes_review_findings():
    base = score(REQUIREMENTS)
    with_review = score(REQUIREMENTS, ReviewResult.model_validate(REVIEW))
    aspects = {a["key"]: a["score"] for a in with_review["aspects"]}
    assert aspects["security"] < {a["key"]: a["score"] for a in base["aspects"]}["security"]
    assert 0 <= with_review["overall"] <= 100
    assert "вспомогательный" in with_review["disclaimer"]


def test_traceability():
    t = traceability(REQUIREMENTS, API, TESTS)
    assert t["uncovered_by_tests"] == ["FR-3"]
    assert t["uncovered_by_ac"] == ["FR-3"]
    assert t["coverage"] == 67


def test_endpoint_conflicts():
    spec = ApiSpec.model_validate(API)
    conflicts = endpoint_conflicts(spec, ["GET /payments/{id} — Get payment", "PUT /refunds"])
    assert conflicts == ["GET /payments/{paymentId} уже есть в проекте: GET /payments/{id} — Get payment"]


# --- рендеринг ---------------------------------------------------------------------------------

def test_sequence_renderers_balance_groups():
    d = SequenceDiagram.model_validate(SEQUENCE)
    mm = sequence_mermaid(d)
    assert mm.startswith("sequenceDiagram")
    assert "alt банк ответил" in mm and "else timeout" in mm and mm.count("end") == 1
    assert "Note over RefundService: retry, статус PENDING" in mm
    pu = sequence_plantuml(d)
    assert pu.startswith("@startuml") and pu.endswith("@enduml")
    assert 'boundary "Bank" as Bank' in pu and "RefundService --> Client: 202 PROCESSING" in pu


def test_unclosed_group_is_closed():
    data = {**SEQUENCE, "steps": SEQUENCE["steps"][:3]}
    assert sequence_mermaid(SequenceDiagram.model_validate(data)).rstrip().endswith("end")


def test_activity_erd_component_renderers():
    from analyst.models import ActivityDiagram, ArchitectureProposal, ERDiagram

    a = ActivityDiagram.model_validate(ACTIVITY)
    assert 'subgraph lane_Refund_Service["Refund Service"]' in activity_mermaid(a)
    assert "d1 -->|нет| e" in activity_mermaid(a)
    assert "state d1 <<choice>>" in activity_plantuml(a)
    e = ERDiagram.model_validate(ERD)
    assert "payment ||--o{ refund" in erd_mermaid(e)
    assert "* id : uuid <<PK>>" in erd_plantuml(e)
    c = component_mermaid(ArchitectureProposal.model_validate(ARCHITECTURE))
    assert "Refund_Service -.->|Kafka| Kafka" in c


def test_openapi_document():
    spec = ApiSpec.model_validate(API)
    doc = openapi_document(spec)
    op = doc["paths"]["/refunds"]["post"]
    body = op["requestBody"]["content"]["application/json"]
    assert body["schema"]["properties"]["targetCard"]["properties"]["id"]["type"] == "string"
    assert body["schema"]["required"] == ["paymentId", "amount"]
    assert body["schema"]["properties"]["amount"]["example"] == 1000
    assert body["example"]["targetCard"]["id"] == "987654"
    assert op["parameters"][0] == {"name": "Idempotency-Key", "in": "header", "required": True,
                                   "description": "UUID", "schema": {"type": "string"}, "example": "7c9e..."}
    assert op["responses"]["409"]["content"]["application/json"]["schema"] == {"$ref": "#/components/schemas/Error"}
    assert doc["paths"]["/payments/{paymentId}"]["get"]["parameters"][0]["in"] == "path"
    assert doc["components"]["securitySchemes"]["auth"]["scheme"] == "bearer"


def test_fields_schema_arrays():
    from analyst.models import ApiField

    f = lambda n, t="string": ApiField(name=n, type=t, format="", required=True, description="", example="",
                                       validation="")
    schema = fields_schema([f("items[].amount", "integer"), f("items[].id")])
    assert schema["properties"]["items"]["type"] == "array"
    assert set(schema["properties"]["items"]["items"]["properties"]) == {"amount", "id"}


def test_artifact_markdown_all_kinds():
    assert "Открытые вопросы" in artifact_markdown("discovery", __import__("analyst_fakes").DISCOVERY_OPEN)
    md = artifact_markdown("requirements", REQUIREMENTS)
    assert "### Functional Requirements" in md and "**Given** платёж проведён" in md
    md = artifact_markdown("review", {"review": REVIEW, "score": score(REQUIREMENTS), "checklist": None})
    assert "Requirements completeness" in md and "⛔ **RV-2" in md
    assert "```mermaid" in artifact_markdown("sequence", SEQUENCE)
    assert "`POST /refunds`" in artifact_markdown("api", {"spec": API, "conflicts": ["x"]})


# --- LLM-обёртка -------------------------------------------------------------------------------

def test_strict_schema_requires_all_fields():
    schema = strict_schema(DiscoveryResult)
    assert schema["additionalProperties"] is False
    assert set(schema["required"]) == set(DiscoveryResult.model_fields)
    question = schema["$defs"]["Question"]
    assert question["additionalProperties"] is False and "suggested_answers" in question["required"]


class _Stream:
    def __init__(self, message):
        self.message = message

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False

    def get_final_message(self):
        return self.message


class _Client:
    def __init__(self, *messages):
        self.queue = list(messages)
        self.requests = []
        self.beta = SimpleNamespace(messages=SimpleNamespace(stream=self._stream))

    def _stream(self, **kwargs):
        self.requests.append(kwargs)
        return _Stream(self.queue.pop(0))


def _msg(text, stop="end_turn"):
    return SimpleNamespace(stop_reason=stop, content=[SimpleNamespace(type="thinking"),
                                                      SimpleNamespace(type="text", text=text)])


def test_claude_llm_request_and_retry_on_invalid_json():
    good = json.dumps({"answer": "ok", "sources": []})
    client = _Client(_msg('{"answer": 1}'), _msg(good))
    llm = ClaudeLLM(client=client, model="claude-opus-5-5")
    from analyst.models import ChatAnswer

    result = llm.structured(system=["agent", "project"], prompt="hi", schema=ChatAnswer, effort="low")
    assert result.answer == "ok"
    req = client.requests[0]
    assert req["model"] == "claude-opus-5-5"
    assert req["thinking"] == {"type": "adaptive"}
    assert req["output_config"]["effort"] == "low"
    assert req["output_config"]["format"]["type"] == "json_schema"
    assert req["fallbacks"] == "default" and req["betas"] == ["server-side-fallback-2026-07-01"]
    assert req["system"][-1]["cache_control"] == {"type": "ephemeral"}
    assert "cache_control" not in req["system"][0]


def test_claude_llm_refusal():
    from analyst.models import ChatAnswer

    llm = ClaudeLLM(client=_Client(_msg("", "refusal")))
    with pytest.raises(LLMRefusal):
        llm.structured(system=["a"], prompt="x", schema=ChatAnswer)


# --- экспорт -----------------------------------------------------------------------------------

def _artifacts():
    arts = {}
    for kind, content in (("requirements", REQUIREMENTS),
                          ("review", {"review": REVIEW, "score": score(REQUIREMENTS), "checklist": None,
                                      "subject": ""}),
                          ("architecture", ARCHITECTURE),
                          ("api", {"spec": API, "openapi": openapi_document(ApiSpec.model_validate(API)),
                                   "conflicts": []}),
                          ("sequence", SEQUENCE), ("test_cases", TESTS)):
        arts[kind] = {"kind": kind, "version": 1, "status": "draft", "content": content}
    return arts


PROJECT = {"name": "Refund"}
TASK = {"id": "t", "title": "Возврат на другую карту", "problem": "Нужен возврат", "stage": "tests"}


@pytest.mark.parametrize("fmt", ["md", "pdf", "docx", "json", "openapi", "mermaid", "plantuml", "jira"])
def test_export_formats(fmt):
    data, media_type, filename = export(fmt, PROJECT, TASK, _artifacts())
    assert data and media_type
    if fmt == "md":
        text = data.decode()
        assert "## Требования — v1" in text and "## Риски" in text and "RISK-1" in text
    if fmt == "pdf":
        assert data.startswith(b"%PDF")
    if fmt == "openapi":
        assert openapi_is_valid(data.decode()) and filename == "openapi.yaml"
        assert yaml.safe_load(data)["paths"]["/refunds"]["post"]["operationId"] == "createRefund"
    if fmt == "jira":
        text = data.decode()
        for section in ("Summary", "Business Requirements", "API Contract", "Acceptance Criteria", "Dependencies",
                        "Technical Notes"):
            assert section in text
    if fmt == "mermaid":
        assert "sequenceDiagram" in data.decode() and "flowchart LR" in data.decode()
